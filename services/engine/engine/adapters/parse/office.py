"""PowerPoint and Word parser adapters (``parse.pptx``, ``parse.docx``; ING-01/ING-03).

PPTX: one slide at a time: the title becomes a level-2 heading ("Slide 3: Title"), body text becomes paragraphs or
bullet lists (indentation kept), tables become Markdown tables, pictures contribute their alt text as captions,
and speaker notes become ``speaker_notes`` blocks. Every block carries its slide number.
DOCX: "Heading N"/"Title" styles become headings, list-styled paragraphs become lists, tables become Markdown.
Office math (OMML) isn't converted here yet; documents that rely on it are routed to Docling in M4-08/09.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from engine.ingest.normalized import Block, NormalizedDoc

PPTX_ID = "parse.pptx"
DOCX_ID = "parse.docx"
VERSION = "1"


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [[c.replace("|", "\\|").replace("\n", " ").strip() for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(norm[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(lines)


def parse_pptx(path: Path, title: str | None = None) -> NormalizedDoc:
    from pptx import Presentation  # imported lazily: heavy module

    prs = Presentation(str(path))
    blocks: list[Block] = []
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        slide_title = (
            title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        )
        blocks.append(
            Block(
                kind="heading",
                level=2,
                slide=number,
                text=f"Slide {number}: {slide_title}" if slide_title else f"Slide {number}",
            )
        )
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            blocks.extend(_shape_blocks(shape, number))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append(Block(kind="speaker_notes", text=notes, slide=number))
    return NormalizedDoc(
        source_kind="pptx",
        title=title,
        blocks=tuple(b for b in blocks if b.text.strip()),
        parser=PPTX_ID,
        parser_version=VERSION,
    )


def _shape_blocks(shape: Any, slide: int) -> list[Block]:
    out: list[Block] = []
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        out.append(Block(kind="table", text=_md_table(rows), slide=slide))
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        bullets: list[str] = []
        plain: list[str] = []
        for p in shape.text_frame.paragraphs:
            text = "".join(r.text for r in p.runs).strip()
            if not text:
                continue
            if p.level > 0 or len(shape.text_frame.paragraphs) > 1:
                bullets.append("  " * p.level + f"- {text}")
            else:
                plain.append(text)
        if plain:
            out.append(Block(kind="paragraph", text=" ".join(plain), slide=slide))
        if bullets:
            out.append(Block(kind="list", text="\n".join(bullets), slide=slide))
    elif shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
        alt = shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr")
        if alt and alt[0].strip():
            out.append(Block(kind="caption", text=f"Figure: {alt[0].strip()}", slide=slide))
    elif getattr(shape, "shapes", None) is not None:  # group shape
        for child in shape.shapes:
            out.extend(_shape_blocks(child, slide))
    return out


def parse_docx(path: Path, title: str | None = None) -> NormalizedDoc:
    import docx  # imported lazily
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    blocks: list[Block] = []
    list_buf: list[str] = []

    def flush_list() -> None:
        if list_buf:
            blocks.append(Block(kind="list", text="\n".join(list_buf)))
            list_buf.clear()

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            flush_list()
            table = Table(child, document)
            blocks.append(Block(kind="table", text=_md_table([[c.text for c in r.cells] for r in table.rows])))
            continue
        if tag != "p":
            continue
        para = Paragraph(child, document)
        text = para.text.strip()
        style = (para.style.name if para.style is not None else "") or ""
        if not text:
            flush_list()
            continue
        if style == "Title":
            flush_list()
            blocks.append(Block(kind="heading", level=1, text=text))
        elif style.startswith("Heading ") and style.split()[-1].isdigit():
            flush_list()
            blocks.append(Block(kind="heading", level=min(6, int(style.split()[-1])), text=text))
        elif "List" in style:
            list_buf.append(f"- {text}")
        else:
            flush_list()
            blocks.append(Block(kind="paragraph", text=text))
    flush_list()
    return NormalizedDoc(source_kind="docx", title=title, blocks=tuple(blocks), parser=DOCX_ID, parser_version=VERSION)
