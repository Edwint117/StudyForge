from pathlib import Path

import docx
from pptx import Presentation
from pptx.util import Inches

from engine.adapters.parse.markdown import parse_markdown, parse_text
from engine.adapters.parse.office import parse_docx, parse_pptx
from engine.ingest.chunking import chunk_document

NOTE = """---
tags: [stats]
---
# Probability

See [[Bayes' theorem|Bayes]] and [[Priors]].

## Conditional probability

**Definition 2.1.** P(A|B) is the probability of A given B.

$$
P(A|B) = \\frac{P(A \\cap B)}{P(B)}
$$

- first item
- second item
  continued

| x | P(x) |
|---|---|
| 1 | 0.5 |

```python
def f(x):
    return x
```
"""


def test_parse_markdown_blocks() -> None:
    doc = parse_markdown(NOTE, title="Stats notes")
    kinds = [b.kind for b in doc.blocks]
    assert kinds == ["heading", "paragraph", "heading", "paragraph", "math", "list", "table", "code"]
    assert doc.blocks[1].text == "See Bayes and Priors."  # wikilinks → plain text, alias preferred
    assert doc.blocks[4].text.startswith("P(A|B) = \\frac")
    assert doc.blocks[5].text.count("\n") == 2  # continuation line kept in the list
    assert doc.blocks[7].language == "python" and "return x" in doc.blocks[7].text
    chunks = chunk_document(doc)
    assert any(c.content_type == "definition" for c in chunks)
    assert chunks[0].heading_path == ("Probability",)


def test_parse_markdown_inline_display_math_and_text() -> None:
    doc = parse_markdown("$$E = mc^2$$\n\nplain")
    assert [(b.kind, b.text) for b in doc.blocks] == [("math", "E = mc^2"), ("paragraph", "plain")]
    text_doc = parse_text("First para\nwraps here.\n\n\nSecond para.")
    assert [b.text for b in text_doc.blocks] == ["First para wraps here.", "Second para."]


def test_parse_pptx(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = "Bayes' theorem"
    body = slide.placeholders[1].text_frame
    body.text = "Relates conditional probabilities"
    p = body.add_paragraph()
    p.text = "Needs P(B) > 0"
    p.level = 1
    slide.notes_slide.notes_text_frame.text = "Mention spam filters."
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide2.shapes.title.text = "Example"
    table = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "x", "P(x)"
    table.cell(1, 0).text, table.cell(1, 1).text = "1", "0.5"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    doc = parse_pptx(path)
    assert doc.source_kind == "pptx"
    headings = [b.text for b in doc.blocks if b.kind == "heading"]
    assert headings == ["Slide 1: Bayes' theorem", "Slide 2: Example"]
    bullets = next(b for b in doc.blocks if b.kind == "list")
    assert bullets.text == "- Relates conditional probabilities\n  - Needs P(B) > 0" and bullets.slide == 1
    notes = next(b for b in doc.blocks if b.kind == "speaker_notes")
    assert notes.text == "Mention spam filters." and notes.slide == 1
    table_block = next(b for b in doc.blocks if b.kind == "table")
    assert table_block.slide == 2 and "| 1 | 0.5 |" in table_block.text
    assert [c.slide for c in chunk_document(doc)] == [1, 2]


def test_parse_docx(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_heading("Lecture 3", level=1)
    d.add_paragraph("Markov chains are memoryless.")
    d.add_heading("Transition matrix", level=2)
    d.add_paragraph("rows sum to one", style="List Bullet")
    d.add_paragraph("entries are probabilities", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "state", "next"
    t.cell(1, 0).text, t.cell(1, 1).text = "A", "B"
    path = tmp_path / "notes.docx"
    d.save(str(path))

    doc = parse_docx(path)
    assert [(b.kind, b.level) for b in doc.blocks] == [
        ("heading", 1),
        ("paragraph", None),
        ("heading", 2),
        ("list", None),
        ("table", None),
    ]
    assert doc.blocks[3].text == "- rows sum to one\n- entries are probabilities"
    assert "| A | B |" in doc.blocks[4].text
    chunks = chunk_document(doc)
    assert chunks[-1].heading_path == ("Lecture 3", "Transition matrix")
